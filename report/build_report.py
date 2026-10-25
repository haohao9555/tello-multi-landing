"""Render the report Markdown to editable DOCX, HTML and browser-generated PDF."""
from pathlib import Path
import re,sys,html,json
sys.path.insert(0,'/tmp/rl-report-deps')
from markdown_it import MarkdownIt
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from playwright.sync_api import sync_playwright

ROOT=Path(__file__).resolve().parent
SOURCE=ROOT/'Project_Report.md'
md=MarkdownIt('commonmark',{'html':False}).enable('table')
tokens=md.parse(SOURCE.read_text())

def slug(s):
    return re.sub(r'[^a-z0-9]+','-',s.lower()).strip('-')

def heading_renderer(tokens,idx,options,env):
    t=tokens[idx];text=tokens[idx+1].content
    attrs=f' id="{slug(text)}"'
    return f'<{t.tag}{attrs}>'
md.renderer.rules['heading_open']=heading_renderer
body=md.render(SOURCE.read_text())
# A proper front-matter TOC for the paginated PDF; page numbers are added after the first pass.
toc=[]
for i,t in enumerate(tokens):
    if t.type=='heading_open' and t.tag in ['h2','h3']:
        text=tokens[i+1].content
        if text.startswith('Assignment') or text in ['Project Report', 'Table of Contents']:continue
        toc.append((text,slug(text),t.tag))
toc_html='<nav class="toc">'+''.join(f'<a class="{tag}" href="#{id}"><span>{html.escape(text)}</span><b data-page="{id}"></b></a>' for text,id,tag in toc)+'</nav>'
body=re.sub(r'(<h2 id="table-of-contents">Table of Contents</h2>).*?(<hr[^>]*>)',r'\1'+toc_html+r'\2',body,flags=re.S)
# Turn image paragraphs into indivisible figures with visible captions.
body=re.sub(r'<p><img src="([^"]+)" alt="([^"]*)"\s*/?></p>',lambda m:f'<figure><img src="{m[1]}" alt="{m[2]}"><figcaption>{m[2]}</figcaption></figure>',body)
body=re.sub(r'<p>(Chenghao Zhuang|Bo Zhao|Junqi Zhang):</p>',r'<p class="signature">\1:</p>',body)
style='''
@page {size:A4;margin:20mm 19mm 19mm}
body{font-family:Georgia,"Times New Roman",serif;font-size:10.7pt;line-height:1.38;color:#17252d}
h1,h2,h3,h4{font-family:Arial,sans-serif;color:#17475c;line-height:1.18;break-after:avoid}
h1{font-size:22pt;margin:27mm 0 10mm}h2{font-size:16pt;margin:8mm 0 4mm}h3{font-size:12pt;margin:6mm 0 3mm}h4{font-size:10.8pt;margin:5mm 0 2mm}
.signature{margin-bottom:10mm;text-align:left;break-inside:avoid}
p{margin:0 0 3mm;orphans:3;widows:3;text-align:justify}
p:has(>strong:only-child){break-after:avoid}
table{border-collapse:collapse;width:100%;font:9pt/1.25 Arial,sans-serif;margin:4mm 0 5mm;break-inside:avoid}
th,td{border:1px solid #c3d1d8;padding:2.1mm 2.3mm;vertical-align:top;overflow-wrap:anywhere}th{background:#eaf1f5;text-align:left;color:#163e52}
tr:nth-child(even) td{background:#f8fafb}thead{display:table-header-group}
figure{margin:5mm 0 6mm;break-inside:avoid}figure img{display:block;width:100%;max-height:175mm;object-fit:contain}
figcaption{font:9pt/1.3 Arial,sans-serif;color:#4c5c65;margin-top:2mm;text-align:left}
hr{border:0;height:0;break-after:page;margin:0}#abstract{margin-top:0}#abstract+p{font-size:10.1pt;line-height:1.28}
[id="1-introduction-and-background"]{break-before:page}
[id="2-2-system-architecture-design"]{break-before:page}
figure:has(img[src="assets/architecture.png"]){break-after:page}
pre{font:8pt/1.35 "DejaVu Sans Mono",monospace;background:#f1f4f6;padding:3mm;white-space:pre-wrap;overflow-wrap:anywhere;break-inside:avoid}
code{font-family:"DejaVu Sans Mono",monospace;font-size:.84em;overflow-wrap:anywhere}
a{color:#175d7b;text-decoration:none}.toc{font:10pt/1.35 Arial,sans-serif}.toc a{display:flex;justify-content:space-between;padding:1mm 0;border-bottom:1px dotted #d4dce0}.toc a.h3{padding-left:5mm;font-size:9.5pt}
'''

def write_html(content):
    target=ROOT/'Project_Report.html'
    target.write_text('<!doctype html><html lang="en"><head><meta charset="utf-8"><title>43008 — Multi-Drone Landing</title><style>'+style+'</style></head><body>'+content+'</body></html>')
    return target

def to_pdf(content):
    target=write_html(content)
    with sync_playwright() as p:
        browser=p.chromium.launch(headless=True,args=['--no-sandbox'])
        page=browser.new_page();page.goto(target.as_uri());page.wait_for_load_state('networkidle')
        page.pdf(path=str(ROOT/'Project_Report.pdf'),format='A4',print_background=True,
                 display_header_footer=True,
                 header_template='<div></div>',
                 footer_template='<div style="font-family:Arial;font-size:8px;width:100%;padding:0 19mm;color:#697882;display:flex;justify-content:space-between"><span></span><span><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>',
                 prefer_css_page_size=True)
        browser.close()

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=Inches(.79);sec.bottom_margin=Inches(.75);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
normal=doc.styles['Normal'];normal.font.name='Cambria';normal.font.size=Pt(10.5)
normal.paragraph_format.line_spacing=1.15;normal.paragraph_format.space_after=Pt(7)
for name,size in [('Title',23),('Heading 1',16),('Heading 2',12),('Heading 3',11)]:
    st=doc.styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('17475C')
    st.paragraph_format.keep_with_next=True
for name in ['Heading 1','Heading 2','Heading 3']:
    doc.styles[name].paragraph_format.space_before=Pt(12)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('').font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
settings=doc.settings.element
update=OxmlElement('w:updateFields');update.set(qn('w:val'),'true');settings.append(update)

def inline(paragraph,children):
    bold=italic=False
    hyperlink=None
    for child in children or []:
        typ=child.type
        if typ=='strong_open':bold=True
        elif typ=='strong_close':bold=False
        elif typ=='em_open':italic=True
        elif typ=='em_close':italic=False
        elif typ=='link_open':
            hyperlink=OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'),paragraph.part.relate_to(child.attrGet('href'),RT.HYPERLINK,is_external=True))
            paragraph._p.append(hyperlink)
        elif typ=='link_close':hyperlink=None
        elif typ in ['text','code_inline']:
            r=paragraph.add_run(child.content);r.bold=bold;r.italic=italic
            if typ=='code_inline':r.font.name='Consolas';r.font.size=Pt(8.5)
            if hyperlink is not None:
                r.font.color.rgb=RGBColor.from_string('175D7B');r.underline=True
                hyperlink.append(r._r)
        elif typ in ['softbreak','hardbreak']:paragraph.add_run().add_break()
        elif typ=='image':
            image=ROOT/child.attrGet('src')
            from PIL import Image
            with Image.open(image) as im:w,h=im.size
            width=min(6.65,6.7*w/h)
            paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next=True
            paragraph.add_run().add_picture(str(image),width=Inches(width))
            caption=doc.add_paragraph(child.content);caption.paragraph_format.space_after=Pt(10)
            for run in caption.runs:run.font.size=Pt(9);run.italic=True
            if child.attrGet('src')=='assets/architecture.png':doc.add_page_break()

i=0;cover=True
while i<len(tokens):
    t=tokens[i]
    if t.type=='heading_open':
        text=tokens[i+1].content
        if text=='Table of Contents':
            doc.add_heading(text,level=1)
            p=doc.add_paragraph()
            start=OxmlElement('w:fldChar');start.set(qn('w:fldCharType'),'begin');p.add_run()._r.append(start)
            instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' TOC \\o "1-2" \\h \\z \\u '
            p.add_run()._r.append(instr)
            sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');p.add_run()._r.append(sep)
            for label,_,tag in toc:
                p.add_run(('    ' if tag=='h3' else '')+label+'\n')
            end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');p.add_run()._r.append(end)
            i+=3
            while i<len(tokens) and tokens[i].type!='hr':i+=1
            continue
        if text in ['1. Introduction and Background','2.2 System Architecture Design']:doc.add_page_break()
        if cover and (text.startswith('Assignment') or text == 'Project Report'):
            p=doc.add_paragraph(style='Subtitle')
        else:
            p=doc.add_heading(level=max(1,int(t.tag[1:])-1)) if t.tag!='h1' else doc.add_paragraph(style='Title')
        inline(p,tokens[i+1].children);i+=3;continue
    if t.type=='hr':doc.add_page_break();cover=False;i+=1;continue
    if t.type=='paragraph_open':
        p=doc.add_paragraph()
        if tokens[i+1].content.startswith('**Table '):p.paragraph_format.keep_with_next=True
        inline(p,tokens[i+1].children)
        if tokens[i+1].content in ['Chenghao Zhuang:', 'Bo Zhao:', 'Junqi Zhang:']:p.paragraph_format.space_after=Pt(25)
        i+=3;continue
    if t.type=='fence':
        p=doc.add_paragraph();r=p.add_run(t.content.rstrip());r.font.name='Consolas';r.font.size=Pt(8)
        p.paragraph_format.keep_together=True;i+=1;continue
    if t.type=='table_open':
        rows=[];current=None;i+=1
        while i<len(tokens) and tokens[i].type!='table_close':
            z=tokens[i]
            if z.type=='tr_open':current=[]
            elif z.type=='inline':current.append(z)
            elif z.type=='tr_close':rows.append(current)
            i+=1
        table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
        for ridx,row in enumerate(rows):
            cells=table.add_row().cells
            for cell,token in zip(cells,row):
                p=cell.paragraphs[0];inline(p,token.children)
                for run in p.runs:run.font.size=Pt(9);run.bold=ridx==0
            if ridx==0:
                prop=table.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');prop.append(rep)
            prop=table.rows[-1]._tr.get_or_add_trPr();avoid=OxmlElement('w:cantSplit');prop.append(avoid)
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
        i+=1;continue
    i+=1
doc.save(ROOT/'Project_Report.docx')
to_pdf(body)
# Add PDF page references to the contents after the first pagination pass.
try:
    import pymupdf
    pdf=pymupdf.open(ROOT/'Project_Report.pdf')
    pages=[page.get_text() for page in pdf]
    for label,id,_ in toc:
        # TOC is page 2; prefer body headings after it.
        found=[n+1 for n,page in enumerate(pages) if n>=2 and label in page]
        if found:body=body.replace(f'<b data-page="{id}"></b>',f'<b data-page="{id}">{found[0]}</b>')
    pdf.close()
    to_pdf(body)
except ImportError:
    pass
print(json.dumps({'docx':str(ROOT/'Project_Report.docx'),'pdf':str(ROOT/'Project_Report.pdf'),'words':len(SOURCE.read_text().split())}))
