"""Build an editable Word report with native Office Math equations."""
from pathlib import Path
import copy,subprocess,zipfile,sys
sys.path.insert(0,'/tmp/rl-report-deps')
from docx import Document
from docx.shared import Pt
R=Path(__file__).resolve().parent
reference=R/'Project_Report.docx';old=Document(reference)
# Retain the existing cover and contents field. Pandoc supplies the revised body.
prefix=[]
for node in old.element.body:
 if node.tag.endswith('sectPr'):break
 text=''.join(node.itertext())
 if getattr(node,'tag','').endswith('}p') and any(p._p is node and p.text=='Abstract' for p in old.paragraphs):break
 prefix.append(copy.deepcopy(node))
target=R/'Project_Report_math.tmp.docx'
subprocess.run(['pandoc',str(R/'overleaf/word_math.md'),'-f','markdown+raw_tex+tex_math_dollars','-t','docx','--reference-doc='+str(reference),'--resource-path='+str(R),'-o',str(target)],check=True)
doc=Document(target)
for node in reversed(prefix):doc.element.body.insert(0,node)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.12
for p in doc.paragraphs:
 if p.text in ['1. Introduction and Background','2.2 System Architecture Design','3. Results and Evaluation','4. Discussion and Conclusions','5. References','Appendices']:p.paragraph_format.page_break_before=True
 if p.text in ['Chenghao Zhuang:','Bo Zhao:','Junqi Zhang:']:p.paragraph_format.space_after=Pt(25)
for section in doc.sections:
 for footer in [section.footer,section.first_page_footer,section.even_page_footer]:
  for p in footer.paragraphs:
   for run in p.runs:
    if 'Part E' in run.text:run.text=''
doc.core_properties.title='Robust Autonomous Multi-UAV Landing under Wind Disturbances Using PPO'
doc.save(target)
with zipfile.ZipFile(target) as z:
 xml=z.read('word/document.xml').decode();assert xml.count('<m:oMathPara')>=8,'Native equations missing'
assert len(Document(target).inline_shapes)==23
assert any('autonomous landing of four UAVs' in p.text for p in Document(target).paragraphs)
target.replace(reference)
print('Editable Word report updated with native Office Math and 23 figures.')

# Publish an offline HTML version with native MathML, using the same formulas.
subprocess.run(['pandoc',str(R/'overleaf/web_math.md'),'-f','markdown+raw_tex+tex_math_dollars','-t','html5','--standalone','--mathml','--metadata','pagetitle=Robust Autonomous Multi-UAV Landing under Wind Disturbances Using PPO','-o',str(R/'Project_Report.html')],check=True)
p=R/'Project_Report.html';html=p.read_text();style='<style>body{max-width:1000px;margin:3rem auto;padding:0 2rem;font:18px/1.55 Georgia,serif;color:#17252d}h1,h2,h3{font-family:Arial,sans-serif;color:#17475c;line-height:1.25}figure{margin:2rem 0}figure img{display:block;max-width:100%;max-height:700px;margin:auto}figcaption{font-size:.9em;color:#4c5c65}table{border-collapse:collapse;width:100%;margin:1.4rem 0;font-size:.9em}td,th{border:1px solid #c3d1d8;padding:.5rem;text-align:left}th{background:#eaf1f5}a{color:#175d7b}math[display=block]{overflow-x:auto;padding:1rem 0}@media print{body{font-size:11pt;margin:0;padding:0}figure,table{break-inside:avoid}}</style>'
p.write_text(html.replace('</head>',style+'</head>'))
print('Offline HTML updated with native MathML equations.')
